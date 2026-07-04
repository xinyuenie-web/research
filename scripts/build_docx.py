# -*- coding: utf-8 -*-
"""将 docs/法治与人治.md 构建为排版规范的 Word 文档。"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC = Path(__file__).resolve().parent.parent / "docs" / "法治与人治.md"
OUT = Path(__file__).resolve().parent.parent / "docs" / "法治与人治.docx"

HEADING_COLOR = RGBColor(0x1F, 0x3B, 0x63)


def set_font(run, name_east_asia="宋体", name_latin="Times New Roman",
             size=Pt(12), bold=False, color=None):
    run.font.name = name_latin
    run.font.size = size
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name_east_asia)


def add_para(doc, text, *, east_asia="宋体", size=Pt(12), bold=False,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_indent=True,
             space_after=Pt(6), color=None):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = space_after
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    if first_line_indent:
        pf.first_line_indent = Pt(24)  # 首行缩进两字符
    run = p.add_run(text)
    set_font(run, name_east_asia=east_asia, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.4
    if level == 1:
        pf.space_before, pf.space_after = Pt(18), Pt(10)
        size, ea = Pt(16), "黑体"
    else:
        pf.space_before, pf.space_after = Pt(12), Pt(6)
        size, ea = Pt(14), "黑体"
    run = p.add_run(text)
    set_font(run, name_east_asia=ea, size=size, bold=True, color=HEADING_COLOR)
    # 标题样式便于生成目录/导航
    p.style = doc.styles["Heading %d" % level]
    for r in p.runs:
        set_font(r, name_east_asia=ea, size=size, bold=True, color=HEADING_COLOR)
    return p


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = table.cell(i, j)
            cell.paragraphs[0].text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(cell_text)
            set_font(run, size=Pt(10.5), bold=(i == 0))
    doc.add_paragraph()


def cjk_count(text):
    return len(re.findall(r"[\u4e00-\u9fff\u3000-\u303f\uff00-\uffef]", text))


def main():
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()

    for section in doc.sections:
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(2.54)
        section.left_margin = section.right_margin = Cm(3.0)

    body_text = []
    in_table, table_rows = False, []
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        i += 1
        if not line.strip():
            continue
        if line == "@table":
            in_table, table_rows = True, []
            continue
        if line == "@endtable":
            add_table(doc, table_rows)
            in_table = False
            continue
        if in_table:
            cells = [c.strip() for c in line.strip("|").split("|")]
            table_rows.append(cells)
            body_text.append("".join(cells))
            continue
        if line.startswith("@subtitle "):
            add_para(doc, line[len("@subtitle "):], size=Pt(14), bold=True,
                     align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=False,
                     space_after=Pt(4))
            continue
        if line.startswith("@meta "):
            add_para(doc, line[len("@meta "):], size=Pt(11),
                     align=WD_ALIGN_PARAGRAPH.CENTER, first_line_indent=False,
                     space_after=Pt(18), color=RGBColor(0x66, 0x66, 0x66))
            continue
        if line.startswith("# "):
            title = line[2:]
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(title)
            set_font(run, name_east_asia="黑体", size=Pt(24), bold=True,
                     color=HEADING_COLOR)
            continue
        if line.startswith("## "):
            add_heading(doc, line[3:], 1)
            body_text.append(line[3:])
            continue
        if line.startswith("### "):
            add_heading(doc, line[4:], 2)
            body_text.append(line[4:])
            continue
        add_para(doc, line)
        body_text.append(line)

    doc.save(OUT)
    total = cjk_count("".join(body_text))
    print("Saved:", OUT)
    print("正文中文字符数(含标点,不含大标题/副题):", total)
    if total < 5000:
        print("WARNING: 少于 5000 字!", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
